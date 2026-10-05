#!/usr/bin/env python3
"""Build a provenance-aware workbook of saved central-charge prefactors.

This is a read-only reduction of existing repository artifacts.  It does not
rerun circuit dynamics.  The generated workbook keeps the fitted prefactor,
its source units, the conversion to natural-log units, and the central-charge
normalization explicit so full-strip and single-wall estimators are not mixed.
"""

from __future__ import annotations

import json
import math
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference, ScatterChart, Series
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def find_repo_root(start: Path) -> Path:
    for candidate in (start, *start.parents):
        if (candidate / "PROJECT_ADMIN" / "REPO_POLICY.md").exists():
            return candidate
    raise RuntimeError("Could not locate repository root")


SCRIPT_PATH = Path(__file__).resolve()
ROOT = find_repo_root(SCRIPT_PATH.parent)
OUT_DIR = SCRIPT_PATH.parent
WORKBOOK_PATH = OUT_DIR / "central_charge_prefactor_inventory_v1.xlsx"
MASTER_CSV_PATH = OUT_DIR / "central_charge_prefactor_master.csv"
MANIFEST_PATH = OUT_DIR / "workbook_manifest.json"

LN2 = math.log(2.0)
TARGET_FULL = 1.0 / 3.0
TARGET_WALL = 1.0 / 6.0

MASTER_COLUMNS = [
    "record_id",
    "provenance_tier",
    "evidence_status",
    "campaign_family",
    "observable",
    "estimator",
    "protocol",
    "postselect_probability",
    "Nx",
    "Ny",
    "cycles",
    "cycle",
    "sample_count",
    "sample_note",
    "n_shell",
    "alpha_top",
    "alpha_triv",
    "trial_orbital",
    "dw_truncation",
    "measure_slab_only",
    "wall_type",
    "wall_location",
    "initial_state",
    "sequence",
    "y0_averaged",
    "subsystem",
    "wall_sector",
    "fit_window",
    "n_fit_points",
    "source_slope",
    "source_slope_units",
    "prefactor_m_nats",
    "prefactor_stderr_nats",
    "target_m_nats",
    "deviation_pct",
    "c_factor",
    "c_eff_normalized",
    "c_total_contribution",
    "c_sem",
    "r2",
    "r2_status",
    "source_kind",
    "source_path",
    "notes",
]

PROVENANCE_ORDER = {
    "A_canonical_saved": 0,
    "A_exact_control": 1,
    "B_saved_reconstructible": 2,
    "C_audit_refit": 3,
    "D_weak_legacy": 4,
    "E_stale_excluded": 5,
}


def rel(path: Path | str) -> str:
    p = Path(path)
    try:
        return str(p.resolve().relative_to(ROOT))
    except ValueError:
        return str(p)


def finite_or_blank(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        value = float(value)
    if isinstance(value, float) and not math.isfinite(value):
        return ""
    if isinstance(value, (list, tuple, dict)):
        return json.dumps(value, sort_keys=True)
    return value


def linear_fit(x: np.ndarray, y: np.ndarray) -> dict[str, float]:
    x = np.asarray(x, dtype=float)
    y = np.asarray(y, dtype=float)
    mask = np.isfinite(x) & np.isfinite(y)
    x, y = x[mask], y[mask]
    if len(x) < 3:
        raise ValueError("Need at least three finite points for a fit")
    coeff, cov = np.polyfit(x, y, 1, cov=True)
    slope, intercept = float(coeff[0]), float(coeff[1])
    stderr = float(math.sqrt(max(float(cov[0, 0]), 0.0)))
    pred = slope * x + intercept
    ss_res = float(np.sum((y - pred) ** 2))
    ss_tot = float(np.sum((y - np.mean(y)) ** 2))
    r2 = 1.0 - ss_res / ss_tot if ss_tot > 0 else 1.0
    return {
        "slope": slope,
        "intercept": intercept,
        "stderr": stderr,
        "r2": r2,
        "n": int(len(x)),
    }


def record(**kwargs: Any) -> dict[str, Any]:
    row = {column: "" for column in MASTER_COLUMNS}
    row.update(kwargs)
    m = row.get("prefactor_m_nats")
    target = row.get("target_m_nats")
    factor = row.get("c_factor")
    m_se = row.get("prefactor_stderr_nats")
    if m != "" and m is not None:
        m = float(m)
        row["prefactor_m_nats"] = m
        if row.get("c_total_contribution", "") in ("", None):
            row["c_total_contribution"] = 3.0 * m
        if factor not in ("", None) and row.get("c_eff_normalized", "") in ("", None):
            row["c_eff_normalized"] = float(factor) * m
        if target not in ("", None) and row.get("deviation_pct", "") in ("", None):
            row["deviation_pct"] = 100.0 * (m / float(target) - 1.0)
        if m_se not in ("", None) and factor not in ("", None) and row.get("c_sem", "") in ("", None):
            row["c_sem"] = float(factor) * float(m_se)
    if row.get("r2") not in ("", None):
        row["r2_status"] = row.get("r2_status") or "available"
    else:
        row["r2_status"] = row.get("r2_status") or "not reported in surviving source"
    return {key: finite_or_blank(row.get(key, "")) for key in MASTER_COLUMNS}


def fit_accumulator(path: Path) -> pd.DataFrame:
    with np.load(path, allow_pickle=False) as z:
        ay = np.asarray(z["ay_values"], dtype=int)
        cycles = np.asarray(z["cycles"], dtype=int)
        ny = int(np.asarray(z["ny"]).item())
        rows: list[dict[str, Any]] = []
        fit_mask = ay >= 8
        x = np.log((ny / np.pi) * np.sin(np.pi * ay[fit_mask] / ny))
        for index, cycle in enumerate(cycles):
            means = []
            counts = []
            for a in ay:
                key = f"Ay{a:03d}"
                count = float(np.asarray(z[f"count_{key}"])[index])
                total = float(np.asarray(z[f"total_sum_{key}"])[index])
                counts.append(count)
                means.append(total / count if count > 0 else np.nan)
            fit = linear_fit(x, np.asarray(means)[fit_mask])
            rows.append(
                {
                    "cycle": int(cycle),
                    "sample_count": int(round(min(counts) / ny)),
                    "slope": fit["slope"],
                    "slope_err": fit["stderr"],
                    "r2": fit["r2"],
                    "n_fit_points": fit["n"],
                    "Ay_fit_min": int(ay[fit_mask].min()),
                    "Ay_fit_max": int(ay[fit_mask].max()),
                }
            )
    return pd.DataFrame(rows)


def collect_data() -> tuple[pd.DataFrame, dict[str, pd.DataFrame], pd.DataFrame]:
    master: list[dict[str, Any]] = []
    sheets: dict[str, pd.DataFrame] = {}
    record_counter = 0

    def add(**kwargs: Any) -> None:
        nonlocal record_counter
        record_counter += 1
        kwargs.setdefault("record_id", f"C-{record_counter:05d}")
        master.append(record(**kwargs))

    # 1. High-statistics per-trajectory and ensemble PC/postselection archive.
    highstat_path = ROOT / (
        "00_WORKSPACE/COLAB/colab_charge_fluctuations/analysis_outputs/"
        "streaming_covariance_scaling_fits/"
        "N20_selected_nsh1_dwtrunc1_init-default_S100_cycles-2Ny/"
        "tables/streaming_covariance_scaling_fit_summary.csv"
    )
    traj = pd.read_csv(highstat_path)
    traj["c_eff"] = 3.0 * traj["entropy_slope"]
    traj["target_prefactor_nats"] = TARGET_FULL
    traj["prefactor_deviation_pct"] = 100.0 * (traj["entropy_slope"] / TARGET_FULL - 1.0)
    sheets["Trajectory_Fits"] = traj
    grouped = (
        traj.groupby(["config_id", "protocol", "Nx", "Ny", "cycle_label"], as_index=False)
        .agg(
            sample_count=("sample_index", "count"),
            prefactor_m_nats=("entropy_slope", "mean"),
            prefactor_std_nats=("entropy_slope", "std"),
            r2=("entropy_r2", "mean"),
        )
    )
    grouped["prefactor_sem_nats"] = grouped["prefactor_std_nats"] / np.sqrt(grouped["sample_count"])
    grouped["c_eff"] = 3.0 * grouped["prefactor_m_nats"]
    grouped["c_sem"] = 3.0 * grouped["prefactor_sem_nats"]
    grouped["is_final_cycle"] = grouped["cycle_label"] == 2 * grouped["Ny"]
    sheets["HighStat_PC_PS"] = grouped
    for row in grouped.itertuples(index=False):
        protocol = str(row.protocol)
        add(
            provenance_tier="A_canonical_saved",
            evidence_status="canonical ensemble mean; SEM derived from saved trajectory fits",
            campaign_family="highstat_PC_PS_2Ny",
            observable="EE",
            estimator="full-x y0-averaged strip log-chord",
            protocol=protocol,
            Nx=int(row.Nx),
            Ny=int(row.Ny),
            cycles=2 * int(row.Ny),
            cycle=int(row.cycle_label),
            sample_count=int(row.sample_count),
            sample_note="100 independent trajectories" if protocol == "perfect_correction" else "one deterministic postselected record",
            n_shell=1,
            alpha_top=1.0,
            alpha_triv=30.0,
            trial_orbital="X",
            dw_truncation=True,
            measure_slab_only="inferred True from bundled default",
            wall_type="hard/support-truncated",
            wall_location="[5,15]",
            initial_state="default pure",
            sequence="raster_y",
            y0_averaged=True,
            subsystem="full x; y strip",
            wall_sector="both walls",
            fit_window=f"Ay=8..{int(row.Ny)//2}",
            n_fit_points=int(row.Ny) // 2 - 7,
            source_slope=float(row.prefactor_m_nats),
            source_slope_units="natural entropy",
            prefactor_m_nats=float(row.prefactor_m_nats),
            prefactor_stderr_nats=float(row.prefactor_sem_nats) if np.isfinite(row.prefactor_sem_nats) else "",
            target_m_nats=TARGET_FULL,
            c_factor=3.0,
            r2=float(row.r2),
            source_kind="saved per-trajectory fit table; ensemble grouping",
            source_path=rel(highstat_path),
            notes="Postselection sample count is actual, not the requested campaign count.",
        )

    # 2. Dedicated large-system fixed-cycle and time-scaled size sweeps.
    large_root = ROOT / "00_WORKSPACE/COLAB/colab_large_entanglement_scaling_N20/gpu_data"
    size_rows: list[pd.DataFrame] = []
    size_specs = [
        ("pure_state_entanglement_slope_vs_system_size/runs", "fixed_C50_S100", 100),
        ("pure_state_entanglement_slope_vs_system_size_time_scaling/runs", "time_scaled_C_Ny_over_2_S10", 10),
    ]
    for relative_dir, family, samples in size_specs:
        for path in sorted((large_root / relative_dir).glob("N20x*/fit_rows.csv")):
            df = pd.read_csv(path)
            df["campaign_family"] = family
            df["samples"] = samples
            df["source_path"] = rel(path)
            size_rows.append(df)
            for row in df.itertuples(index=False):
                add(
                    provenance_tier="A_canonical_saved",
                    evidence_status="completed saved fit",
                    campaign_family=family,
                    observable="EE",
                    estimator="full-x y0-averaged strip log-chord",
                    protocol="perfect_correction",
                    Nx=int(row.Nx),
                    Ny=int(row.Ny),
                    cycles=int(row.cycle),
                    cycle=int(row.cycle),
                    sample_count=samples,
                    sample_note=f"{samples} trajectories",
                    n_shell=int(row.nshell) if pd.notna(row.nshell) else "",
                    alpha_top=1.0,
                    alpha_triv=30.0,
                    trial_orbital="X",
                    dw_truncation=True,
                    measure_slab_only="inferred True from bundled default",
                    wall_type="hard/support-truncated",
                    wall_location="[5,15]",
                    initial_state="default pure",
                    sequence="raster_y",
                    y0_averaged=True,
                    subsystem="full x; y strip",
                    wall_sector="both walls",
                    fit_window=f"Ay={int(row.Ay_fit_min)}..{int(row.Ay_fit_max)}",
                    n_fit_points=int(row.n_fit_points),
                    source_slope=float(row.slope),
                    source_slope_units="natural entropy",
                    prefactor_m_nats=float(row.slope),
                    prefactor_stderr_nats=float(row.slope_err),
                    target_m_nats=TARGET_FULL,
                    c_factor=3.0,
                    r2=float(row.r2),
                    source_kind="saved fit_rows.csv",
                    source_path=rel(path),
                )
    sheets["Size_Sweeps"] = pd.concat(size_rows, ignore_index=True)

    # Recover incomplete Ny=60 fixed-C run.
    ny60_acc = large_root / (
        "pure_state_entanglement_slope_vs_system_size/runs/"
        "N20x60_nsh1_dwtrunc1_C50_S100/accumulator_checkpoint.npz"
    )
    recovered_rows: list[pd.DataFrame] = []
    if ny60_acc.exists():
        rec = fit_accumulator(ny60_acc)
        rec["source"] = rel(ny60_acc)
        rec["campaign_family"] = "fixed_C50_S100_incomplete"
        rec["Nx"] = 20
        rec["Ny"] = 60
        recovered_rows.append(rec)
        for row in rec.itertuples(index=False):
            add(
                provenance_tier="C_audit_refit",
                evidence_status="incomplete accumulator; 80/100 samples",
                campaign_family="fixed_C50_S100_incomplete",
                observable="EE",
                estimator="full-x y0-averaged strip log-chord",
                protocol="perfect_correction",
                Nx=20,
                Ny=60,
                cycles=50,
                cycle=int(row.cycle),
                sample_count=int(row.sample_count),
                sample_note="80 of 100 requested trajectories recovered from accumulator",
                n_shell=1,
                alpha_top=1.0,
                alpha_triv=30.0,
                trial_orbital="X",
                dw_truncation=True,
                measure_slab_only="inferred True from bundled default",
                wall_type="hard/support-truncated",
                wall_location="[5,15]",
                initial_state="default pure",
                sequence="raster_y",
                y0_averaged=True,
                subsystem="full x; y strip",
                wall_sector="both walls",
                fit_window=f"Ay={int(row.Ay_fit_min)}..{int(row.Ay_fit_max)}",
                n_fit_points=int(row.n_fit_points),
                source_slope=float(row.slope),
                source_slope_units="natural entropy",
                prefactor_m_nats=float(row.slope),
                prefactor_stderr_nats=float(row.slope_err),
                target_m_nats=TARGET_FULL,
                c_factor=3.0,
                r2=float(row.r2),
                source_kind="audit reconstruction from accumulator sums/counts",
                source_path=rel(ny60_acc),
            )

    # 3. Complete S=10 cycle sweep and recoverable S=100/S=90 accumulator.
    cycle_path = large_root / (
        "pure_state_entanglement_slope_vs_cycle/runs/"
        "N20x40_nsh1_dwtrunc1_C100_S10/slope_vs_cycle.csv"
    )
    cycle10 = pd.read_csv(cycle_path)
    cycle10["c_eff"] = 3.0 * cycle10["slope"]
    cycle10["samples"] = 10
    cycle10["evidence_status"] = "completed saved fit"
    cycle10["source_path"] = rel(cycle_path)
    for row in cycle10.itertuples(index=False):
        add(
            provenance_tier="A_canonical_saved",
            evidence_status="completed saved cycle fit",
            campaign_family="N20x40_cycle_S10",
            observable="EE",
            estimator="full-x y0-averaged strip log-chord",
            protocol="perfect_correction",
            Nx=20,
            Ny=40,
            cycles=100,
            cycle=int(row.cycle),
            sample_count=10,
            sample_note="10 trajectories",
            n_shell=1,
            alpha_top=1.0,
            alpha_triv=30.0,
            trial_orbital="X",
            dw_truncation=True,
            measure_slab_only="inferred True from bundled default",
            wall_type="hard/support-truncated",
            wall_location="[5,15]",
            initial_state="default pure",
            sequence="raster_y",
            y0_averaged=True,
            subsystem="full x; y strip",
            wall_sector="both walls",
            fit_window=f"Ay={int(row.Ay_fit_min)}..{int(row.Ay_fit_max)}",
            n_fit_points=int(row.n_fit_points),
            source_slope=float(row.slope),
            source_slope_units="natural entropy",
            prefactor_m_nats=float(row.slope),
            prefactor_stderr_nats=float(row.slope_err),
            target_m_nats=TARGET_FULL,
            c_factor=3.0,
            r2=float(row.r2),
            source_kind="saved slope_vs_cycle.csv",
            source_path=rel(cycle_path),
        )
    cycle_acc = large_root / (
        "pure_state_entanglement_slope_vs_cycle/runs/"
        "N20x40_nsh1_dwtrunc1_C100_S100/accumulator_checkpoint.npz"
    )
    if cycle_acc.exists():
        rec = fit_accumulator(cycle_acc)
        rec["source"] = rel(cycle_acc)
        rec["campaign_family"] = "N20x40_cycle_S100_incomplete"
        rec["Nx"] = 20
        rec["Ny"] = 40
        recovered_rows.append(rec)
        for row in rec.itertuples(index=False):
            add(
                provenance_tier="C_audit_refit",
                evidence_status="accumulator only; S=100 through t20 and S=90 afterward",
                campaign_family="N20x40_cycle_S100_incomplete",
                observable="EE",
                estimator="full-x y0-averaged strip log-chord",
                protocol="perfect_correction",
                Nx=20,
                Ny=40,
                cycles=100,
                cycle=int(row.cycle),
                sample_count=int(row.sample_count),
                sample_note="100 trajectories through t20; 90 thereafter",
                n_shell=1,
                alpha_top=1.0,
                alpha_triv=30.0,
                trial_orbital="X",
                dw_truncation=True,
                measure_slab_only="inferred True from bundled default",
                wall_type="hard/support-truncated",
                wall_location="[5,15]",
                initial_state="default pure",
                sequence="raster_y",
                y0_averaged=True,
                subsystem="full x; y strip",
                wall_sector="both walls",
                fit_window=f"Ay={int(row.Ay_fit_min)}..{int(row.Ay_fit_max)}",
                n_fit_points=int(row.n_fit_points),
                source_slope=float(row.slope),
                source_slope_units="natural entropy",
                prefactor_m_nats=float(row.slope),
                prefactor_stderr_nats=float(row.slope_err),
                target_m_nats=TARGET_FULL,
                c_factor=3.0,
                r2=float(row.r2),
                source_kind="audit reconstruction from accumulator sums/counts",
                source_path=rel(cycle_acc),
            )
    sheets["Cycle_Sweeps"] = cycle10
    sheets["Recovered_Accumulators"] = pd.concat(recovered_rows, ignore_index=True)

    # 4. Earlier compact PC/postselection protocol characterization.
    proto_path = ROOT / (
        "00_WORKSPACE/COLAB/colab_small_system_testing/gpu_data/"
        "streaming_covariance_protocol_characterization/campaigns/"
        "N20_multi_geometry_C50_dwtrunc1_init-default_nsh1_S10/ensemble_summaries.csv"
    )
    proto = pd.read_csv(proto_path)
    proto = proto[(proto["metric"] == "entropy_slope") & (proto["group_kind"] == "cycle_snapshot")].copy()
    proto_sample_path = proto_path.parent / "per_sample_cycle_metrics.csv"
    proto_samples = pd.read_csv(proto_sample_path)
    proto_r2 = proto_samples.groupby(
        ["config_id", "protocol", "Nx", "Ny", "cycle_label"], as_index=False
    ).agg(
        entropy_r2_mean=("entropy_r2", "mean"),
        entropy_r2_min=("entropy_r2", "min"),
        entropy_r2_max=("entropy_r2", "max"),
    )
    proto = proto.merge(
        proto_r2,
        on=["config_id", "protocol", "Nx", "Ny", "cycle_label"],
        how="left",
        validate="one_to_one",
    )
    proto["prefactor_sem_nats"] = proto["std"] / np.sqrt(proto["count"])
    proto["c_eff"] = 3.0 * proto["mean"]
    sheets["Early_PC_PS"] = proto
    for row in proto.itertuples(index=False):
        add(
            provenance_tier="A_canonical_saved",
            evidence_status="completed protocol-characterization snapshot",
            campaign_family="early_PC_PS_C50",
            observable="EE",
            estimator="full-x y0-averaged strip log-chord",
            protocol=str(row.protocol),
            Nx=int(row.Nx),
            Ny=int(row.Ny),
            cycles=50,
            cycle=int(row.cycle_label),
            sample_count=int(row.count),
            sample_note="10 trajectories" if int(row.count) > 1 else "one deterministic postselected record",
            n_shell=1,
            alpha_top=1.0,
            alpha_triv=30.0,
            trial_orbital="X",
            dw_truncation=True,
            measure_slab_only="inferred True from bundled default",
            wall_type="hard/support-truncated",
            wall_location="source run metadata",
            initial_state="default pure",
            sequence="raster_y",
            y0_averaged=True,
            subsystem="full x; y strip",
            wall_sector="both walls",
            fit_window="campaign default log-chord window",
            source_slope=float(row.mean),
            source_slope_units="natural entropy",
            prefactor_m_nats=float(row.mean),
            prefactor_stderr_nats=float(row.prefactor_sem_nats) if np.isfinite(row.prefactor_sem_nats) else "",
            target_m_nats=TARGET_FULL,
            c_factor=3.0,
            r2=float(row.entropy_r2_mean),
            r2_status="mean of saved per-sample entropy_r2 values",
            source_kind="saved ensemble_summaries.csv plus per_sample_cycle_metrics.csv",
            source_path=rel(proto_path),
            notes=f"R2 range across records: {float(row.entropy_r2_min):.8g}..{float(row.entropy_r2_max):.8g}; R2 source={rel(proto_sample_path)}",
        )

    # 5. Saved broad-window contours, saved late full-x fits, and new late wall refits.
    contour_root = ROOT / (
        "00_WORKSPACE/COLAB/colab_small_system_testing/analysis_outputs/"
        "pure_state_entanglement_vs_system_size_cpu"
    )
    broad_path = contour_root / "log_chord_fit_rows.csv"
    broad = pd.read_csv(broad_path)
    broad["fit_window_kind"] = "broad_saved"
    late_full_path = contour_root / "full_x_late_window_log_chord_fit_rows.csv"
    late_full = pd.read_csv(late_full_path)
    late_full["fit_window_kind"] = "late_saved"
    contour_frames = [broad.copy(), late_full.copy()]
    for frame, source_path, status, family in [
        (broad, broad_path, "completed saved contour fit", "wall_integrated_broad_saved"),
        (late_full, late_full_path, "completed saved late-window full-x fit", "wall_integrated_late_full_saved"),
    ]:
        for row in frame.itertuples(index=False):
            is_wall = "wall" in str(row.x_interval_kind)
            m_nats = float(row.slope_over_ln2) * LN2
            m_se = float(row.slope_err_over_ln2) * LN2
            add(
                provenance_tier="A_canonical_saved",
                evidence_status=status,
                campaign_family=family,
                observable="EE contour integral",
                estimator="x-window integrated entropy contour vs log chord",
                protocol="perfect_correction",
                Nx=int(row.Nx),
                Ny=int(row.Ny),
                cycles=50,
                cycle=int(row.cycle),
                sample_count=10,
                sample_note="10 trajectories x every y0; processed contour curve",
                n_shell=int(row.nshell) if pd.notna(row.nshell) else "",
                alpha_top=1.0,
                alpha_triv=30.0,
                trial_orbital="X",
                dw_truncation=True,
                measure_slab_only="inferred True from bundled default",
                wall_type="hard/support-truncated",
                wall_location="N16 source metadata",
                initial_state="default pure",
                sequence="raster_y",
                y0_averaged=True,
                subsystem="integrated x-window of full-strip contour",
                wall_sector=str(row.x_interval_kind),
                fit_window=f"Ay={int(row.Ay_fit_min)}..{int(row.Ay_fit_max)}",
                n_fit_points=int(row.n_fit_points),
                source_slope=float(row.slope_over_ln2),
                source_slope_units="S/ln2 (bits)",
                prefactor_m_nats=m_nats,
                prefactor_stderr_nats=m_se,
                target_m_nats=TARGET_WALL if is_wall else TARGET_FULL,
                c_factor=6.0 if is_wall else 3.0,
                r2=float(row.r2),
                source_kind="saved processed contour fit",
                source_path=rel(source_path),
                notes="c_eff_normalized uses 6m for a single wall and 3m for the full strip.",
            )
    curves_path = contour_root / "integrated_contour_curves.csv"
    curves = pd.read_csv(curves_path)
    late_wall_rows: list[dict[str, Any]] = []
    q = curves[(curves["cycle"] == 50) & (curves["x_interval_kind"] != "full_x") & (curves["Ay"] >= 8)]
    for keys, group in q.groupby(["case_id", "Nx", "Ny", "nshell", "x_interval_kind", "x_interval_start", "x_interval_stop"]):
        case_id, nx, ny, nshell, sector, x0, x1 = keys
        fit = linear_fit(
            group["log_sin_pi_Ay_over_Ny"].to_numpy(),
            group["integrated_entropy_mean_over_ln2"].to_numpy(),
        )
        m_nats = fit["slope"] * LN2
        m_se = fit["stderr"] * LN2
        late_wall_rows.append(
            {
                "case_id": case_id,
                "Nx": nx,
                "Ny": ny,
                "nshell": nshell,
                "cycle": 50,
                "x_interval_kind": sector,
                "x_interval_start": x0,
                "x_interval_stop": x1,
                "Ay_fit_min": int(group["Ay"].min()),
                "Ay_fit_max": int(group["Ay"].max()),
                "n_fit_points": fit["n"],
                "slope_over_ln2": fit["slope"],
                "slope_err_over_ln2": fit["stderr"],
                "prefactor_m_nats": m_nats,
                "c_wall": 6.0 * m_nats,
                "r2": fit["r2"],
                "source_path": rel(curves_path),
            }
        )
        add(
            provenance_tier="C_audit_refit",
            evidence_status="new late-window refit of saved integrated curve",
            campaign_family="wall_integrated_late_refit",
            observable="EE contour integral",
            estimator="x-window integrated entropy contour vs log chord",
            protocol="perfect_correction",
            Nx=int(nx),
            Ny=int(ny),
            cycles=50,
            cycle=50,
            sample_count=10,
            sample_note="10 trajectories x every y0; mean contour curve; cross-bin covariance unavailable",
            n_shell=int(nshell),
            alpha_top=1.0,
            alpha_triv=30.0,
            trial_orbital="X",
            dw_truncation=True,
            measure_slab_only="inferred True from bundled default",
            wall_type="hard/support-truncated",
            wall_location="N16 source metadata",
            initial_state="default pure",
            sequence="raster_y",
            y0_averaged=True,
            subsystem=f"integrated x={int(x0)}..{int(x1)} contour window",
            wall_sector=str(sector),
            fit_window=f"Ay={int(group['Ay'].min())}..{int(group['Ay'].max())}",
            n_fit_points=fit["n"],
            source_slope=fit["slope"],
            source_slope_units="S/ln2 (bits)",
            prefactor_m_nats=m_nats,
            prefactor_stderr_nats=m_se,
            target_m_nats=TARGET_WALL,
            c_factor=6.0,
            r2=fit["r2"],
            source_kind="audit refit from saved integrated_contour_curves.csv",
            source_path=rel(curves_path),
            notes="Uses the same Ay>=8 window as the precision full-strip fit.",
        )
    late_wall = pd.DataFrame(late_wall_rows)
    contour_frames.append(late_wall)
    sheets["Wall_EE_Contours"] = pd.concat(contour_frames, ignore_index=True, sort=False)

    # 6. Partial-postselection final-cycle refits.
    partial_root = ROOT / (
        "00_WORKSPACE/COLAB/colab_partial_post-select/gpu_data/"
        "streaming_covariance_observables/campaigns/"
        "N20x40_nsh1_dwtrunc1_alpha2-30_init-default_S10_C40_partial-postselect-psweep/runs"
    )
    partial_rows: list[dict[str, Any]] = []
    for run_dir in partial_root.iterdir():
        entropy_path = run_dir / "entropy_y0avg_vs_ay.npz"
        if not entropy_path.exists():
            continue
        match = re.search(r"psp([0-9.]+)", run_dir.name)
        if not match:
            continue
        psp = float(match.group(1))
        with np.load(entropy_path, allow_pickle=False) as z:
            ay = np.asarray(z["ay_values"], dtype=int)
            values = np.asarray(z["entropy_y0avg_vs_ay"], dtype=float)[:, -1, :]
        mask = ay >= 8
        x = np.log((40.0 / np.pi) * np.sin(np.pi * ay[mask] / 40.0))
        fits = [linear_fit(x, sample[mask]) for sample in values]
        slopes = np.asarray([fit["slope"] for fit in fits])
        slope_mean = float(np.mean(slopes))
        slope_sem = float(np.std(slopes, ddof=1) / np.sqrt(len(slopes))) if len(slopes) > 1 else np.nan
        r2_mean = float(np.mean([fit["r2"] for fit in fits]))
        partial_rows.append(
            {
                "postselect_probability": psp,
                "sample_count": len(slopes),
                "prefactor_m_nats": slope_mean,
                "prefactor_sem_nats": slope_sem,
                "c_eff": 3.0 * slope_mean,
                "c_sem": 3.0 * slope_sem if np.isfinite(slope_sem) else np.nan,
                "r2_mean": r2_mean,
                "source_path": rel(entropy_path),
            }
        )
        add(
            provenance_tier="C_audit_refit",
            evidence_status="new final-cycle fit of saved compact entropy arrays",
            campaign_family="partial_postselection_probability_sweep",
            observable="EE",
            estimator="full-x y0-averaged strip log-chord",
            protocol="postselect" if psp == 1.0 else "partial_postselection",
            postselect_probability=psp,
            Nx=20,
            Ny=40,
            cycles=40,
            cycle=40,
            sample_count=len(slopes),
            sample_note="one deterministic record" if len(slopes) == 1 else f"{len(slopes)} trajectories",
            n_shell=1,
            alpha_top=1.0,
            alpha_triv=30.0,
            trial_orbital="X",
            dw_truncation=True,
            measure_slab_only="inferred True from bundled default",
            wall_type="hard/support-truncated",
            wall_location="[5,15]",
            initial_state="default pure",
            sequence="raster_y",
            y0_averaged=True,
            subsystem="full x; y strip",
            wall_sector="both walls",
            fit_window="Ay=8..20",
            n_fit_points=13,
            source_slope=slope_mean,
            source_slope_units="natural entropy",
            prefactor_m_nats=slope_mean,
            prefactor_stderr_nats=slope_sem if np.isfinite(slope_sem) else "",
            target_m_nats=TARGET_FULL,
            c_factor=3.0,
            r2=r2_mean,
            source_kind="audit refit from saved entropy_y0avg_vs_ay.npz",
            source_path=rel(entropy_path),
        )
    sheets["Partial_Postselect"] = pd.DataFrame(partial_rows).sort_values("postselect_probability")

    # 7. Exact MI/CMI calibration and historical nonperfect stochastic reconstruction.
    mi_json_path = ROOT / (
        "00_WORKSPACE/CURRENT/experiment_review/legacy_evidence_figure_atlas/data/"
        "result_05_tripartite_information_reduction.json"
    )
    mi_payload = json.loads(mi_json_path.read_text())
    mi_rows = pd.DataFrame(mi_payload["fit_rows"])
    mi_rows["c_eff"] = mi_rows["c_factor"] * mi_rows["slope"]
    mi_rows["source_path"] = rel(mi_json_path)
    sheets["Exact_MI_CMI"] = mi_rows
    for row in mi_rows.itertuples(index=False):
        is_wall = "wall" in str(row.window)
        exact = str(row.evidence) == "exact"
        add(
            provenance_tier="A_exact_control" if exact else "D_weak_legacy",
            evidence_status="exact deterministic calibration" if exact else "reconstructed obsolete nonperfect-feedback pilot",
            campaign_family="exact_MI_CMI_calibration" if exact else "legacy_stochastic_MI_CMI",
            observable="CMI / contour-window contribution",
            estimator="I2-I3 vs z=log(1/(1-x))",
            protocol="exact_ground_state" if exact else "nonperfect_stochastic",
            Nx=20 if exact else 12,
            Ny=40 if exact else 31,
            cycles="" if exact else 100,
            cycle="" if exact else 100,
            sample_count=1 if exact else 100,
            sample_note="deterministic" if exact else "100 trajectories; whole-trajectory bootstrap available",
            n_shell="" if not exact else "exact source-era construction",
            alpha_top=1.0 if exact else "",
            alpha_triv=30.0 if exact else "",
            dw_truncation="not applicable",
            wall_type="exact domain wall" if exact else "legacy domain wall",
            wall_location="[6,14] source-era exact" if exact else "legacy N12",
            initial_state="ground state" if exact else "default pure",
            sequence="" if exact else "dw_symmetric",
            y0_averaged=False if exact else True,
            subsystem="full-x covariance; optional post-contour x window",
            wall_sector=str(row.window),
            fit_window="all stored z rows",
            source_slope=float(row.slope),
            source_slope_units="natural entropy",
            prefactor_m_nats=float(row.slope),
            prefactor_stderr_nats=float(row.slope_error),
            target_m_nats=float(row.target),
            c_factor=float(row.c_factor),
            r2=float(row.r2),
            source_kind="saved atlas JSON/NPZ",
            source_path=rel(mi_json_path),
            notes="Only full width is literally a CMI; wall rows are post-contour contributions.",
        )

    # 8. Deterministic flattened-H and exact-domain-wall EE controls.
    flat_path = ROOT / "notebooks/flattened_hamiltonian_analysis/data/slope_fit_df.csv"
    flat = pd.read_csv(flat_path)
    flat["prefactor_m_nats"] = flat["slope"]
    flat["c_factor_normalized"] = np.where(flat["x_interval_kind"] == "wall_window", 6.0, 3.0)
    flat["c_eff_normalized"] = flat["prefactor_m_nats"] * flat["c_factor_normalized"]
    flat["c_total_contribution"] = 3.0 * flat["prefactor_m_nats"]
    sheets["FlatH_Exact_Controls"] = flat
    for row in flat.itertuples(index=False):
        is_wall = str(row.x_interval_kind) == "wall_window"
        add(
            provenance_tier="A_exact_control",
            evidence_status="deterministic flattened-H control",
            campaign_family="flattened_H_exact_control",
            observable="EE",
            estimator="flattened-H strip or wall-contour log-chord",
            protocol="exact_flattened_ground_state",
            Nx=int(row.Nx),
            Ny=int(row.Ny),
            sample_count=1,
            sample_note="deterministic",
            n_shell=int(row.nshell) if pd.notna(row.nshell) else "",
            alpha_top=float(row.alpha_1),
            alpha_triv=float(row.alpha_2),
            trial_orbital=str(row.trial_orbital),
            dw_truncation=bool(row.dw_truncation),
            measure_slab_only="not trajectory dynamics",
            wall_type="hard/support-truncated" if bool(row.dw_truncation) else "soft/untruncated",
            wall_location="[5,15] windows in saved table",
            initial_state=str(row.state_type),
            y0_averaged="not recorded in this summary table",
            subsystem=str(row.x_interval),
            wall_sector=str(row.x_interval_kind),
            fit_window=str(row.branch),
            source_slope=float(row.slope),
            source_slope_units="natural entropy",
            prefactor_m_nats=float(row.slope),
            prefactor_stderr_nats=float(row.slope_err),
            target_m_nats=TARGET_WALL if is_wall else TARGET_FULL,
            c_factor=6.0 if is_wall else 3.0,
            r2=float(row.r2),
            source_kind="saved flattened-H fit CSV",
            source_path=rel(flat_path),
            notes=f"configuration={row.configuration_label}; branch={row.branch}",
        )
    exact_ee_path = ROOT / (
        "00_WORKSPACE/CURRENT/experiment_review/b0_exact_domain_wall/results/"
        "20260816_191957/processed/tables/entropy_fits.csv"
    )
    exact_ee = pd.read_csv(exact_ee_path)
    exact_ee["c_eff_normalized"] = np.where(
        exact_ee["quantity"] == "physical_wall_contour",
        2.0 * exact_ee["c_estimate"],
        exact_ee["c_estimate"],
    )
    sheets["Exact_DW_EE_Controls"] = exact_ee
    for row in exact_ee.itertuples(index=False):
        is_wall = str(row.quantity) == "physical_wall_contour"
        add(
            provenance_tier="A_exact_control",
            evidence_status="deterministic exact-domain-wall control",
            campaign_family="b0_exact_domain_wall",
            observable=f"Renyi-q={row.q} EE",
            estimator=str(row.quantity),
            protocol="exact_ground_state",
            Nx=int(row.nx),
            Ny=int(row.ny),
            sample_count=1,
            sample_note="deterministic",
            wall_type=str(row.construction),
            initial_state="exact ground state",
            y0_averaged="source-specific exact geometry",
            subsystem="physical wall contour" if is_wall else "full strip",
            wall_sector=f"wall {int(row.wall_index)}" if is_wall else "both walls",
            fit_window=f"Ay={int(row.ay_min)}..{int(row.ay_max)}",
            n_fit_points=int(row.n_points),
            source_slope=float(row.slope),
            source_slope_units="natural Renyi entropy",
            prefactor_m_nats=float(row.slope),
            target_m_nats="",
            c_factor="q-dependent",
            c_eff_normalized=float(row.c_estimate) * (2.0 if is_wall else 1.0),
            c_total_contribution=float(row.c_estimate),
            r2=float(row.r2),
            source_kind="saved exact-domain-wall entropy_fits.csv",
            source_path=rel(exact_ee_path),
            notes="Source c_estimate is the total-central-charge contribution; a single wall is doubled in c_eff_normalized.",
        )

    # 9. Weak and stale legacy rows retained for provenance, never mixed with canonical claims.
    legacy_rows: list[dict[str, Any]] = []
    haining_path = ROOT / "notebooks/slope_vs_system_size_from_plots_Haining_Data.ipynb"
    haining = {
        "EE": {
            "soft_overlap": [0.5589, 0.4894, 0.4591, 0.4413],
            "hard_truncated": [0.5104, 0.4541, 0.4285, 0.4071],
        },
        "MI_CMI": {
            "soft_overlap": [0.4761, 0.4436, 0.4282, 0.4174],
            "hard_truncated": [0.4272, 0.3993, 0.3942, 0.3804],
        },
    }
    for observable, wall_map in haining.items():
        for wall_label, slopes in wall_map.items():
            for ny, slope in zip([30, 40, 50, 60], slopes):
                legacy_rows.append(
                    {
                        "legacy_family": "Haining_manual_plot_transcription",
                        "observable": observable,
                        "wall_label": wall_label,
                        "Ny": ny,
                        "sample_count": "unknown",
                        "slope": slope,
                        "c_factor": 3.0,
                        "c_eff": 3.0 * slope,
                        "status": "weak provenance; manually transcribed from plot legends",
                        "source_path": rel(haining_path),
                    }
                )
                add(
                    provenance_tier="D_weak_legacy",
                    evidence_status="manual transcription; raw files and exact fit metadata absent",
                    campaign_family="Haining_manual_plot_transcription",
                    observable=observable,
                    estimator="full-width log-chord fit",
                    protocol="perfect_correction_equivalent",
                    Nx=20,
                    Ny=ny,
                    sample_count="unknown",
                    n_shell="unknown",
                    alpha_top="unknown",
                    alpha_triv="unknown",
                    dw_truncation=True if wall_label.startswith("hard") else False,
                    wall_type=wall_label,
                    y0_averaged=True,
                    subsystem="full x",
                    wall_sector="both walls",
                    fit_window="unknown",
                    source_slope=slope,
                    source_slope_units="natural entropy",
                    prefactor_m_nats=slope,
                    target_m_nats=TARGET_FULL,
                    c_factor=3.0,
                    source_kind="manually transcribed notebook values",
                    source_path=rel(haining_path),
                )
    sample_pdf_path = ROOT / "00_WORKSPACE/LEGACY/sample_average_testing"
    prefix_counts = [1, 10, 50, 100]
    prefix_sets = [
        ("LA4_one_wall", [0.2564, 0.2610, 0.2443, 0.2487], 6.0, TARGET_WALL, 0.9997981993676421),
        ("LA4_full_width", [0.5878, 0.5133, 0.5022, 0.5112], 3.0, TARGET_FULL, 0.9999571964498953),
        ("LA3_full_width", [0.5698, 0.5545, 0.5121, 0.5251], 3.0, TARGET_FULL, 0.9998834904080175),
    ]
    for label, slopes, factor, target, s100_r2 in prefix_sets:
        for samples, slope in zip(prefix_counts, slopes):
            legacy_rows.append(
                {
                    "legacy_family": "N12x31_prefix_sample_sweep",
                    "observable": "MI_CMI",
                    "wall_label": label,
                    "Ny": 31,
                    "sample_count": samples,
                    "slope": slope,
                    "c_factor": factor,
                    "c_eff": factor * slope,
                    "r2": s100_r2 if samples == 100 else "",
                    "status": "obsolete nonperfect-feedback pilot; PDF-rounded values",
                    "source_path": rel(sample_pdf_path),
                }
            )
            add(
                provenance_tier="D_weak_legacy",
                evidence_status="obsolete nonperfect-feedback pilot; values rounded from saved PDFs",
                campaign_family="N12x31_prefix_sample_sweep",
                observable="MI/CMI contour contribution",
                estimator=label,
                protocol="nonperfect_stochastic",
                Nx=12,
                Ny=31,
                cycles=100,
                cycle=100,
                sample_count=samples,
                sample_note="prefix sample count, not independent reruns",
                wall_type="legacy domain wall",
                initial_state="default pure",
                sequence="dw_symmetric",
                y0_averaged=True,
                subsystem=label,
                wall_sector="one wall" if "one_wall" in label else "both walls",
                fit_window="source PDF definition",
                source_slope=slope,
                source_slope_units="natural entropy",
                prefactor_m_nats=slope,
                target_m_nats=target,
                c_factor=factor,
                r2=s100_r2 if samples == 100 else "",
                r2_status="available from atlas reconstruction" if samples == 100 else "not reported for prefix subset in surviving PDF",
                source_kind="saved legacy PDF / atlas reconstruction",
                source_path=rel(sample_pdf_path),
            )
    stale_notebook = ROOT / "notebooks/analysis.ipynb"
    stale_values = [
        ("perfect_correction", "left", 0.3037125, 250),
        ("perfect_correction", "right", 0.2950453, 250),
        ("nonperfect_adaptive", "left", 0.3597238, 250),
        ("nonperfect_adaptive", "right", 0.3564333, 250),
        ("postselect", "left", 0.2372118, 250),
        ("postselect", "right", 0.2326550, 250),
        ("exact_reduced_strip", "left", 0.2022111, 1),
        ("exact_reduced_strip", "right", 0.2022111, 1),
    ]
    for protocol, sector, slope, samples in stale_values:
        legacy_rows.append(
            {
                "legacy_family": "stale_reduced_strip_MI_notebook",
                "observable": "MI_CMI",
                "wall_label": sector,
                "Ny": 16,
                "sample_count": samples,
                "slope": slope,
                "c_factor": 6.0,
                "c_eff": 6.0 * slope,
                "r2": "",
                "status": "stale/excluded source-output mismatch",
                "source_path": rel(stale_notebook),
            }
        )
        add(
            provenance_tier="E_stale_excluded",
            evidence_status="stale notebook output; live source and embedded output disagree",
            campaign_family="stale_reduced_strip_MI_notebook",
            observable="reduced-strip MI",
            estimator="two-column wall strip",
            protocol=protocol,
            Nx=12,
            Ny=16,
            cycles=20 if protocol != "postselect" else 5,
            sample_count=samples,
            sample_note="embedded output claims S=250; current source caps samples at 2",
            y0_averaged=True,
            subsystem="two-column reduced strip",
            wall_sector=sector,
            fit_window="embedded notebook output",
            source_slope=slope,
            source_slope_units="natural entropy",
            prefactor_m_nats=slope,
            target_m_nats=TARGET_WALL,
            c_factor=6.0,
            r2_status="not reported in embedded notebook output",
            source_kind="embedded notebook output only",
            source_path=rel(stale_notebook),
        )
    sheets["Legacy_Weak_Stale"] = pd.DataFrame(legacy_rows)

    # Dataset index includes raw branches without a validated c reduction.
    dataset_index = pd.DataFrame(
        [
            {
                "dataset": "High-stat hard-wall PC/PS EE",
                "coverage": "Nx20; Ny30/40/50; C=2Ny; PC S100; PS S1; every cycle/y0",
                "prefactor_status": "canonical and included",
                "path": rel(highstat_path),
            },
            {
                "dataset": "Large hard-wall size/cycle EE",
                "coverage": "fixed C50 S100; time-scaled S10; N20x40 C100",
                "prefactor_status": "canonical plus explicitly labeled accumulator refits",
                "path": rel(large_root),
            },
            {
                "dataset": "N16 wall-integrated EE contours",
                "coverage": "Ny30/40; nshell1/2; C5/10/20/50; left/right/full",
                "prefactor_status": "saved broad fits; new late-window wall refits included",
                "path": rel(contour_root),
            },
            {
                "dataset": "Partial-postselection EE",
                "coverage": "N20x40 C40; p=0..1; S10 except p=1 S1",
                "prefactor_status": "raw complete; audit final-cycle refits included",
                "path": rel(partial_root),
            },
            {
                "dataset": "Soft-wall pure-state production bundles",
                "coverage": "N20x20/30; PC; S10/S5; y0 EE and contours",
                "prefactor_status": "not included as a claimed c: no committed stable fit table",
                "path": "00_WORKSPACE/LARGE_RESULTS/classA_final_production_outputs/production_10sample_v4_occupied_frame_cycle_resolved/02_pure_wall_master",
            },
            {
                "dataset": "Hard/soft maximally-mixed purification contours",
                "coverage": "N16x30/32 and N30x40; PC/PS; nshell1/2",
                "prefactor_status": "excluded from spatial c; purification-in-time observable",
                "path": "00_WORKSPACE/COLAB/colab_small_system_testing/gpu_data/purification_entropy_contours_maxmix/campaigns",
            },
            {
                "dataset": "Purification alpha sweep",
                "coverage": "Ny16/24/32; alpha sweep; hard/soft; PC/PS",
                "prefactor_status": "excluded from spatial c; charge-sharpening/purification data",
                "path": "00_WORKSPACE/COLAB/colab_charge_fluctuations/cpu_data/purification_charge_sharpening_alpha_sweep",
            },
            {
                "dataset": "Modern wall-resolved MI",
                "coverage": "canonical PC/PS target protocol",
                "prefactor_status": "not found",
                "path": "",
            },
            {
                "dataset": "Haining hard/soft slopes",
                "coverage": "Nx20; Ny30/40/50/60; EE and MI",
                "prefactor_status": "included only as weak manual transcription",
                "path": rel(haining_path),
            },
        ]
    )

    master_df = pd.DataFrame(master, columns=MASTER_COLUMNS)
    master_df["_prov_order"] = master_df["provenance_tier"].map(PROVENANCE_ORDER).fillna(99)
    master_df = master_df.sort_values(
        ["_prov_order", "campaign_family", "observable", "Nx", "Ny", "cycle", "wall_sector"],
        kind="stable",
    ).drop(columns="_prov_order")
    return master_df, sheets, dataset_index


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
THIN_GRAY = Side(style="thin", color="D9E1F2")
TIER_FILLS = {
    "A_canonical_saved": "E2F0D9",
    "A_exact_control": "DDEBF7",
    "B_saved_reconstructible": "FFF2CC",
    "C_audit_refit": "FCE4D6",
    "D_weak_legacy": "F4CCCC",
    "E_stale_excluded": "D9D9D9",
}


def safe_sheet_name(name: str) -> str:
    return re.sub(r"[\\/*?:\[\]]", "_", name)[:31]


def write_dataframe(ws, df: pd.DataFrame, table_name: str, freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    headers = list(df.columns)
    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_index, value=str(header))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)
    for row_index, row in enumerate(df.itertuples(index=False, name=None), start=2):
        for col_index, value in enumerate(row, start=1):
            cell = ws.cell(row=row_index, column=col_index, value=finite_or_blank(value))
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if headers[col_index - 1] in {"source_path", "path"} and cell.value:
                target = ROOT / str(cell.value)
                cell.hyperlink = target.as_uri()
                cell.style = "Hyperlink"
    if len(df) > 0:
        ref = f"A1:{get_column_letter(len(headers))}{len(df) + 1}"
        table = Table(displayName=re.sub(r"[^A-Za-z0-9_]", "_", table_name)[:250], ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False
        )
        ws.add_table(table)
        ws.auto_filter.ref = ref
    for col_index, header in enumerate(headers, start=1):
        sample_values = [str(header)]
        if len(df):
            sample_values.extend(str(v) for v in df.iloc[: min(len(df), 300), col_index - 1].dropna())
        width = min(max(max((len(v) for v in sample_values), default=8) + 2, 10), 42)
        if header in {"notes", "source_path", "path", "coverage", "prefactor_status", "evidence_status"}:
            width = min(max(width, 28), 55)
        ws.column_dimensions[get_column_letter(col_index)].width = width
    ws.row_dimensions[1].height = 34


def add_master_formatting(ws, df: pd.DataFrame) -> None:
    tier_col = list(df.columns).index("provenance_tier") + 1
    c_col = list(df.columns).index("c_eff_normalized") + 1
    for row_idx in range(2, len(df) + 2):
        tier = ws.cell(row=row_idx, column=tier_col).value
        color = TIER_FILLS.get(str(tier))
        if color:
            ws.cell(row=row_idx, column=tier_col).fill = PatternFill("solid", fgColor=color)
    ws.conditional_formatting.add(
        f"{get_column_letter(c_col)}2:{get_column_letter(c_col)}{len(df)+1}",
        ColorScaleRule(start_type="num", start_value=0.8, start_color="F8696B",
                       mid_type="num", mid_value=1.0, mid_color="63BE7B",
                       end_type="num", end_value=1.2, end_color="F8696B"),
    )


def build_dashboard(wb: Workbook, master: pd.DataFrame) -> None:
    ws = wb["Dashboard"]
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Central-charge prefactor inventory"
    ws["A1"].font = Font(size=18, bold=True, color="1F4E78")
    ws["A3"] = "Scope"
    ws["A3"].font = Font(size=12, bold=True)
    ws["A4"] = (
        "All located central-charge/prefactor estimates are normalized with explicit provenance. "
        "Canonical saved results, exact controls, audit refits, weak legacy transcriptions, and stale "
        "outputs remain separate."
    )
    ws.merge_cells("A4:H5")
    ws["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A7"] = "Normalization"
    ws["A7"].font = Font(size=12, bold=True)
    normalization = [
        ("Full strip / full-width CMI", "m_target=1/3", "c=3m"),
        ("Single physical wall", "m_target=1/6", "c_wall=6m"),
        ("Contour slopes stored as S/ln2", "convert m_nats=m_bits*ln2", "then apply 3m or 6m"),
        ("c_total_contribution", "always 3m when applicable", "one wall contributes about 1/2"),
    ]
    for r, values in enumerate(normalization, start=8):
        for c, value in enumerate(values, start=1):
            ws.cell(r, c, value)
            ws.cell(r, c).fill = SUBHEADER_FILL if c == 1 else PatternFill("solid", fgColor="F7FBFF")
            ws.cell(r, c).border = Border(bottom=THIN_GRAY)
    ws["A14"] = "Inventory counts"
    ws["A14"].font = Font(size=12, bold=True)
    counts = master.groupby("provenance_tier").size().reset_index(name="rows")
    for r, row in enumerate(counts.itertuples(index=False), start=15):
        ws.cell(r, 1, row.provenance_tier)
        ws.cell(r, 2, int(row.rows))
        color = TIER_FILLS.get(str(row.provenance_tier))
        if color:
            ws.cell(r, 1).fill = PatternFill("solid", fgColor=color)
    ws["A22"] = "Key current conclusions"
    ws["A22"].font = Font(size=12, bold=True)
    conclusions = [
        "Hard-wall PC full-strip EE converges near c≈1.04–1.08; deterministic postselection is ≈1.002–1.004.",
        "Late-window n_shell=1 integrated wall contours average to the 1/6 prefactor within about 1.3% (Ny30) and 0.1% (Ny40).",
        "Soft-wall trajectory data exist but lack a stable committed central-charge reduction; they are indexed, not promoted to canonical c rows.",
        "No modern canonical wall-resolved PC/PS MI campaign was found.",
    ]
    for r, text in enumerate(conclusions, start=23):
        ws.cell(r, 1, f"• {text}")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        ws.cell(r, 1).alignment = Alignment(wrap_text=True)
    for col, width in {"A": 34, "B": 22, "C": 26, "D": 16, "E": 16, "F": 16, "G": 16, "H": 16}.items():
        ws.column_dimensions[col].width = width

    # Hidden chart data.
    data_ws = wb["_ChartData"]
    final_rows = master[
        (master["campaign_family"] == "highstat_PC_PS_2Ny")
        & (pd.to_numeric(master["cycle"], errors="coerce") == pd.to_numeric(master["cycles"], errors="coerce"))
    ].copy()
    final_pivot = final_rows.pivot_table(index="Ny", columns="protocol", values="c_eff_normalized", aggfunc="first").reset_index()
    data_ws.append(["Ny", "perfect_correction", "postselect"])
    for row in final_pivot.itertuples(index=False):
        data_ws.append(list(row))
    chart = LineChart()
    chart.title = "Final c: perfect correction vs postselection"
    chart.y_axis.title = "c_eff"
    chart.x_axis.title = "Ny"
    chart.height = 7.0
    chart.width = 12.0
    data = Reference(data_ws, min_col=2, max_col=3, min_row=1, max_row=len(final_pivot) + 1)
    cats = Reference(data_ws, min_col=1, min_row=2, max_row=len(final_pivot) + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "J3")

    wall = master[(master["campaign_family"] == "wall_integrated_late_refit") & (master["n_shell"] == 1)].copy()
    wall["series"] = wall["wall_sector"].astype(str)
    wall_pivot = wall.pivot_table(index="Ny", columns="series", values="c_eff_normalized", aggfunc="first").reset_index()
    start = len(final_pivot) + 4
    columns = ["Ny"] + [c for c in wall_pivot.columns if c != "Ny"]
    data_ws.append([])
    for c, value in enumerate(columns, start=1):
        data_ws.cell(start, c, value)
    for r_idx, row in enumerate(wall_pivot[columns].itertuples(index=False, name=None), start=start + 1):
        for c_idx, value in enumerate(row, start=1):
            data_ws.cell(r_idx, c_idx, finite_or_blank(value))
    wall_chart = LineChart()
    wall_chart.title = "Late-window n_shell=1 wall-normalized c"
    wall_chart.y_axis.title = "c_wall=6m"
    wall_chart.x_axis.title = "Ny"
    wall_chart.height = 7.0
    wall_chart.width = 12.0
    data = Reference(data_ws, min_col=2, max_col=len(columns), min_row=start, max_row=start + len(wall_pivot))
    cats = Reference(data_ws, min_col=1, min_row=start + 1, max_row=start + len(wall_pivot))
    wall_chart.add_data(data, titles_from_data=True)
    wall_chart.set_categories(cats)
    ws.add_chart(wall_chart, "J18")
    data_ws.sheet_state = "hidden"


def build_workbook(master: pd.DataFrame, sheets: dict[str, pd.DataFrame], dataset_index: pd.DataFrame) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Dashboard")
    wb.create_sheet("README")
    wb.create_sheet("Master_Catalog")
    for sheet_name in sheets:
        wb.create_sheet(safe_sheet_name(sheet_name))
    wb.create_sheet("Dataset_Index")
    wb.create_sheet("Glossary")
    wb.create_sheet("_ChartData")

    readme = wb["README"]
    readme.sheet_view.showGridLines = False
    readme["A1"] = "How to use this workbook"
    readme["A1"].font = Font(size=18, bold=True, color="1F4E78")
    instructions = [
        ("Generated", datetime.now(timezone.utc).isoformat()),
        ("Repository", str(ROOT)),
        ("Purpose", "Organize every located central-charge/prefactor estimate with explicit protocol, geometry, fit window, units, and provenance."),
        ("Primary table", "Master_Catalog: one normalized row per ensemble/control/legacy estimate."),
        ("Raw detail", "Trajectory_Fits retains all 24,240 saved per-trajectory/per-cycle EE fits."),
        ("Canonical", "A_canonical_saved and A_exact_control rows can be used directly, subject to their stated estimator and finite-size caveats."),
        ("Audit refit", "C_audit_refit is reproducible from saved raw/accumulator data but was not a committed source fit table."),
        ("Weak legacy", "D_weak_legacy lacks full raw provenance or uses an obsolete protocol."),
        ("Excluded", "E_stale_excluded is retained only to prevent accidental reuse."),
        ("Wall convention", "For a single physical wall, c_eff_normalized=6m so the chiral-wall target is 1; c_total_contribution=3m is about 1/2."),
        ("Bit conversion", "Stored contour slopes labeled S/ln2 are multiplied by ln(2) before comparing with 1/6 or 1/3."),
        ("Wall metadata warning", "Modern persisted N20 run metadata uses walls [5,15]; some source-era exact/legacy prose uses [6,14]."),
    ]
    for r, (label, text) in enumerate(instructions, start=3):
        readme.cell(r, 1, label)
        readme.cell(r, 1).font = Font(bold=True)
        readme.cell(r, 1).fill = SUBHEADER_FILL
        readme.cell(r, 2, text)
        readme.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 24
    readme.column_dimensions["B"].width = 115

    master_ws = wb["Master_Catalog"]
    write_dataframe(master_ws, master, "MasterCatalog")
    add_master_formatting(master_ws, master)
    for name, df in sheets.items():
        write_dataframe(wb[safe_sheet_name(name)], df, f"tbl_{safe_sheet_name(name)}")
    write_dataframe(wb["Dataset_Index"], dataset_index, "DatasetIndex")
    glossary = pd.DataFrame(
        [
            ("prefactor_m_nats", "Slope after converting entropy to natural-log units."),
            ("source_slope", "Slope exactly as stored or fitted in the source units."),
            ("target_m_nats", "Expected natural-unit prefactor: normally 1/3 full strip or 1/6 one wall."),
            ("c_eff_normalized", "Central charge normalized to unity for the tested sector: 3m full or 6m one wall."),
            ("c_total_contribution", "3m. A single physical wall therefore contributes approximately 1/2."),
            ("c_sem", "Uncertainty propagated from the slope SEM/fit error using the stated c factor."),
            ("r2", "Coefficient of determination for the stated fit. Blank means the surviving source did not report enough information."),
            ("r2_status", "Explains whether R^2 is available, reconstructed, or absent from the surviving source."),
            ("y0_averaged", "Whether the subsystem origin was averaged around periodic y."),
            ("dw_truncation", "Whether the overlap/Wannier measurement kernel was hard support-truncated."),
            ("measure_slab_only", "Whether measurement updates were restricted to the active domain-wall slab; some legacy values are inferred from defaults."),
            ("PC", "Perfect correction: Born-sampled measurement followed by deterministic correction of an incorrect occupation."),
            ("PS", "Postselection. In the modern saved controls this is one deterministic record, not an independent ensemble."),
        ],
        columns=["field_or_term", "definition"],
    )
    write_dataframe(wb["Glossary"], glossary, "Glossary")
    build_dashboard(wb, master)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(WORKBOOK_PATH)


def validate_workbook(expected_rows: int) -> dict[str, Any]:
    wb = load_workbook(WORKBOOK_PATH, read_only=False, data_only=False)
    required = {
        "Dashboard", "README", "Master_Catalog", "Trajectory_Fits",
        "HighStat_PC_PS", "Wall_EE_Contours", "Exact_MI_CMI", "Fit_Quality",
        "Legacy_Weak_Stale", "Dataset_Index", "Glossary",
    }
    missing = sorted(required.difference(wb.sheetnames))
    if missing:
        raise RuntimeError(f"Workbook missing required sheets: {missing}")
    master_ws = wb["Master_Catalog"]
    actual_rows = master_ws.max_row - 1
    if actual_rows != expected_rows:
        raise RuntimeError(f"Master row mismatch: {actual_rows} != {expected_rows}")
    headers = [cell.value for cell in master_ws[1]]
    for field in ["prefactor_m_nats", "target_m_nats", "c_eff_normalized", "source_path"]:
        if field not in headers:
            raise RuntimeError(f"Missing master field {field}")
    return {
        "workbook": rel(WORKBOOK_PATH),
        "bytes": WORKBOOK_PATH.stat().st_size,
        "sheet_count": len(wb.sheetnames),
        "sheets": wb.sheetnames,
        "master_rows": actual_rows,
        "trajectory_fit_rows": wb["Trajectory_Fits"].max_row - 1,
        "validation": "passed",
    }


def main() -> None:
    master, sheets, dataset_index = collect_data()
    r2_numeric = pd.to_numeric(master["r2"], errors="coerce")
    fit_quality = master.assign(r2_available=r2_numeric.notna()).groupby(
        ["provenance_tier", "campaign_family", "observable"], as_index=False
    ).agg(
        fit_rows=("record_id", "count"),
        r2_available_rows=("r2_available", "sum"),
    )
    fit_quality["r2_missing_rows"] = fit_quality["fit_rows"] - fit_quality["r2_available_rows"]
    fit_quality["r2_coverage_fraction"] = fit_quality["r2_available_rows"] / fit_quality["fit_rows"]
    sheets["Fit_Quality"] = fit_quality
    master.to_csv(MASTER_CSV_PATH, index=False)
    build_workbook(master, sheets, dataset_index)
    manifest = validate_workbook(len(master))
    manifest.update(
        {
            "generated_utc": datetime.now(timezone.utc).isoformat(),
            "master_csv": rel(MASTER_CSV_PATH),
            "master_csv_rows": len(master),
            "provenance_counts": master["provenance_tier"].value_counts().to_dict(),
            "builder": rel(SCRIPT_PATH),
        }
    )
    MANIFEST_PATH.write_text(json.dumps(manifest, indent=2, sort_keys=True) + "\n")
    print(json.dumps(manifest, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
